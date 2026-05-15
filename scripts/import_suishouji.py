#!/usr/bin/env python3
"""
随手记Excel → family_ledger 数据库导入
含分类映射 + 去重（日期+金额+备注）
"""
import openpyxl
import mysql.connector
from datetime import datetime

DB_CONFIG = {
    'host': '127.0.0.1', 'port': 3306,
    'user': 'hermes', 'password': 'hermes_ledger_2026',
    'database': 'family_ledger', 'charset': 'utf8mb4',
}

EXCEL_PATH = '/home/vonlohengramm/.hermes/cache/documents/doc_d71d16390f82_随手记家庭账本120260515.xlsx'

# 随手记 → 系统分类ID
CAT_MAP = {
    # 食品酒水
    ('食品酒水', '伙食费'): 110,    # 餐饮·外卖/在外就餐
    ('食品酒水', '饮料酒水'): 113,  # 餐饮·饮料酒水
    ('食品酒水', '买菜'): 111,      # 餐饮·超市买菜
    ('食品酒水', '零食'): 112,      # 餐饮·零食
    ('食品酒水', '中餐'): 110,      # 餐饮·外卖/在外就餐
    ('食品酒水', '水果'): 112,      # 餐饮·零食
    ('食品酒水', '早餐'): 110,      # 餐饮·外卖/在外就餐
    ('食品酒水', '晚餐'): 110,      # 餐饮·外卖/在外就餐
    # 行车交通
    ('行车交通', '停车'): 103,      # 交通·停车费
    ('行车交通', '充电'): 101,      # 交通·充电
    ('行车交通', '自行车'): 106,    # 交通·自行车
    ('行车交通', '加油'): 102,      # 交通·加油
    ('行车交通', '打车'): 105,      # 交通·打车
    ('行车交通', '地铁'): 104,      # 交通·公交地铁
    ('行车交通', '违章罚款'): 108,  # 交通·违章罚款
    ('行车交通', '保险'): 107,      # 交通·保险
    ('行车交通', '维修'): 163,      # 汽车·保养
    ('行车交通', '洗车'): 163,      # 汽车·保养
    # 宝宝费用
    ('宝宝费用', '宝宝用品'): 133,  # 母婴·用品
    ('宝宝费用', '医疗护理'): 130,  # 母婴·医疗
    ('宝宝费用', '妈妈用品'): 133,  # 母婴·用品
    ('宝宝费用', '宝宝食品'): 133,  # 母婴·用品
    ('宝宝费用', '宝宝其他'): 133,  # 母婴·用品
    # 人情费用
    ('人情费用', '孝敬长辈'): 156,  # 人情·孝敬长辈
    ('人情费用', '礼物'): 157,      # 人情·礼物
    ('人情费用', '请客'): 157,      # 人情·礼物
    ('人情费用', '红包'): 155,      # 人情·红包（出）
    # 购物消费
    ('购物消费', '衣裤鞋帽'): 115,  # 购物·服饰
    ('购物消费', '宠物支出'): 135,  # 宠物·猫粮
    ('购物消费', '美妆护肤'): 116,  # 购物·美容
    ('购物消费', '日常用品'): 149,  # 日用·其他日用
    ('购物消费', '家居饰品'): 119,  # 购物·其他购物
    ('购物消费', '电子数码'): 143,  # 数码·硬件
    ('购物消费', '办公用品'): 172,  # 办公·工作支出
    ('购物消费', '清洁用品'): 147,  # 日用·清洁用品
    ('购物消费', '家用纺织'): 119,  # 购物·其他购物
    ('购物消费', '家具家电'): 127,  # 家居·家用电器
    ('购物消费', '厨房用品'): 149,  # 日用·其他日用
    ('购物消费', '珠宝首饰'): 117,  # 购物·饰品
    ('购物消费', '洗护用品'): 147,  # 日用·清洁用品
    ('购物消费', '书报杂志'): 114,  # 购物·图书
    # 休闲娱乐
    ('休闲娱乐', '网游'): 150,      # 游戏·充值
    ('休闲娱乐', '其他娱乐'): 154,  # 娱乐·玩乐
    ('休闲娱乐', '温泉洗浴'): 154,  # 娱乐·玩乐
    ('休闲娱乐', '彩票'): 154,      # 娱乐·玩乐
    ('休闲娱乐', '话剧'): 154,      # 娱乐·玩乐
    ('休闲娱乐', '电影'): 152,      # 娱乐·电影
    # 医疗教育
    ('医疗教育', '药品费'): 141,    # 医疗·药费
    ('医疗教育', '治疗费'): 139,    # 医疗·就诊
    # 金融保险
    ('金融保险', '房贷'): 173,      # 金融·还款
    ('金融保险', '税费'): 175,      # 金融·保险
    ('金融保险', '人身保险'): 175,  # 金融·保险
    # 居家生活
    ('居家生活', '水费'): 121,      # 家居·水费
    ('居家生活', '快递费'): 168,    # 快递·邮费
    ('居家生活', '维修费'): 176,    # 维修·家电维修
    ('居家生活', '燃气费'): 122,    # 家居·煤气
    ('居家生活', '理发'): 116,      # 购物·美容
    ('居家生活', '电费'): 120,      # 家居·电费
    ('居家生活', '物业费'): 126,    # 家居·物业
    # 出差旅游
    ('出差旅游', '娱乐费'): 167,    # 旅行·门票
    # 交流通讯
    ('交流通讯', '手机话费'): 123,  # 家居·话费
    ('交流通讯', '网费'): 124,      # 家居·网费
    # 其他
    ('其他杂项', '其他支出'): 195,  # 其他·其他
    # 装修
    ('装修费用', '家电家具'): 128,  # 家居·装修
}

MEMBER_MAP = {
    '爸爸👨': 1,
    '妈妈👩': 2,
}

def main():
    wb = openpyxl.load_workbook(EXCEL_PATH)
    ws = wb.active

    # === 构建现有数据去重索引 ===
    conn = mysql.connector.connect(**DB_CONFIG)
    c = conn.cursor()
    c.execute("SELECT DATE(trans_date), ABS(amount), IFNULL(note,'') FROM transaction")
    existing = set()
    for dt, amt, note in c.fetchall():
        existing.add((str(dt), float(amt), note))

    # === 解析Excel ===
    new_records = []
    skipped_cat = {}
    skipped_dup = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or row[0] is None:
            continue
        
        tx_type = row[0].strip()       # 支出
        date_str = row[1].strip()       # '2026-05-15 10:09:29'
        cat1 = row[2].strip() if row[2] else ''
        cat2 = row[3].strip() if row[3] else ''
        amount = float(row[7])
        member = row[8].strip() if row[8] else '爸爸👨'
        note = (row[13] or '').strip() if len(row) > 13 else ''

        # 日期解析
        dt = datetime.strptime(date_str, '%Y-%m-%d %H:%M:%S')
        trans_date = dt.date()
        trans_time = dt.time()

        # 金额：支出为负
        amount_val = -abs(amount)

        # 分类映射
        cat_key = (cat1, cat2)
        if cat_key not in CAT_MAP:
            skipped_cat[str(cat_key)] = skipped_cat.get(str(cat_key), 0) + 1
            continue
        category_id = CAT_MAP[cat_key]

        # 成员映射
        user_id = MEMBER_MAP.get(member, 1)

        # 去重：日期+绝对值+备注
        dedup_key = (str(trans_date), abs(amount_val), note)
        if dedup_key in existing:
            skipped_dup += 1
            continue

        new_records.append((user_id, category_id, amount_val, trans_date, trans_time, note[:500] if note else None))

    # === 批量插入 ===
    ins_sql = """INSERT INTO transaction (user_id, category_id, amount, trans_date, trans_time, note, created_at)
                 VALUES (%s, %s, %s, %s, %s, %s, NOW())"""
    
    batch_size = 500
    inserted = 0
    for i in range(0, len(new_records), batch_size):
        batch = new_records[i:i+batch_size]
        c.executemany(ins_sql, batch)
        conn.commit()
        inserted += len(batch)
        print(f"  已导入 {inserted}/{len(new_records)} ...")

    c.execute("SELECT COUNT(*) FROM transaction")
    total = c.fetchone()[0]
    c.close()
    conn.close()

    print(f"\n======= 导入完成 =======")
    print(f"  Excel总行数: {ws.max_row - 1}")
    print(f"  新导入:     {inserted} 笔")
    print(f"  重复跳过:   {skipped_dup} 笔")
    if skipped_cat:
        print(f"  未知分类:   {sum(skipped_cat.values())} 笔")
        for k, v in skipped_cat.items():
            print(f"    {k}: {v}")
    print(f"  数据库总计: {total} 笔")

if __name__ == '__main__':
    main()
